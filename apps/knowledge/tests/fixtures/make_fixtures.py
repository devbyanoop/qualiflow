"""Generates real .docx and .xlsx test fixtures for Phase 1 step 1's
parsing/chunking/delta-hash tests -- not hand-crafted binary blobs,
actual files built via the same libraries real users' documents come
from (python-docx, openpyxl), so Docling parses them the same way it
would parse a real upload.

Run once: `python tests/fixtures/make_fixtures.py` from apps/knowledge.
Produces brd_v1.docx and brd_v2.docx (one heading's content changed,
proving chunk-level delta-hash), and a small coverage_matrix.xlsx.
"""

from pathlib import Path

from docx import Document
from openpyxl import Workbook

FIXTURES_DIR = Path(__file__).parent


def make_brd_v1() -> None:
    doc = Document()
    doc.add_heading("Account Opening Requirements", level=1)
    doc.add_paragraph("This document describes the rules for opening a new account.")

    doc.add_heading("Minimum Age", level=2)
    doc.add_paragraph("A customer must be at least 18 years old to open an account.")

    doc.add_heading("Identity Verification", level=2)
    doc.add_paragraph("A government-issued photo ID is required before an account can be activated.")

    doc.save(FIXTURES_DIR / "brd_v1.docx")


def make_brd_v2() -> None:
    """Same as v1, except the Minimum Age section's wording changed --
    the delta-hash proof is that only this section's chunk re-embeds,
    while Identity Verification's chunk (byte-identical) is reused."""
    doc = Document()
    doc.add_heading("Account Opening Requirements", level=1)
    doc.add_paragraph("This document describes the rules for opening a new account.")

    doc.add_heading("Minimum Age", level=2)
    doc.add_paragraph(
        "A customer must be at least 21 years old to open an account, raised from 18 per the updated policy."
    )

    doc.add_heading("Identity Verification", level=2)
    doc.add_paragraph("A government-issued photo ID is required before an account can be activated.")

    doc.save(FIXTURES_DIR / "brd_v2.docx")


def make_coverage_matrix_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage"
    ws.append(["Requirement", "Covered By", "Status"])
    ws.append(["Minimum age check", "TC-001", "Pass"])
    ws.append(["Identity verification", "TC-002", "Pass"])
    wb.save(FIXTURES_DIR / "coverage_matrix.xlsx")


if __name__ == "__main__":
    FIXTURES_DIR.mkdir(exist_ok=True)
    make_brd_v1()
    make_brd_v2()
    make_coverage_matrix_xlsx()
    print(f"Fixtures written to {FIXTURES_DIR}")
